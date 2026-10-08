"""Export du dossier de réponse : dossier HTML et documents de travail.

Le CLI ne produit pas le mémoire Word : l'agent de l'utilisateur le fait avec
ses propres outils (`remporte guide export`). `dossier.html` est la page
autonome à deux onglets d'origine. Trois documents de travail complètent l'export, sur la charte
Remporte (remporte.charte) : `analyse.html` (analyse et go/no-go mis en page
pour un dirigeant), `feuille-de-route.html` (avancement, sections, trous à
combler) et `matrice-conformite.xlsx` (exigences du CRT ou sections du plan).
"""

from __future__ import annotations

import html
import json
import re
from pathlib import Path

from remporte import cadre, charte, espace


def exporter(dossier: Path, formats: set[str]) -> dict:
    """Exporte les formats demandés.

    Rend un dict {html, analyse, feuille, matrice} → chemin ou None. Le
    dossier demande un plan avec sections ; les documents de
    travail se contentent des fichiers d'analyse déjà remplis.
    """
    dossier = Path(dossier)
    if not (dossier / ".remporte").is_dir():
        raise espace.Erreur(
            f"{dossier} n'est pas un dossier de réponse (pas de .remporte/)."
        )
    sections = espace.sections_plan(dossier)
    if "html" in formats and not sections:
        raise espace.Erreur(
            "04-plan.md ne liste aucune section au format `- [ ] NN — Titre`. "
            "Complétez le plan (remporte guide plan) avant d'exporter."
        )
    resultat: dict[str, Path | None] = {
        "html": None, "analyse": None, "feuille": None,
        "matrice": None,
    }
    if "html" in formats:
        resultat["html"] = _exporter_html(dossier, sections)
    # La matrice passe avant l'analyse, qui renvoie vers elle avec son bilan.
    bilan_matrice = None
    if "matrice" in formats:
        resultat["matrice"], bilan_matrice = _exporter_matrice(dossier, sections)
    if "analyse" in formats:
        resultat["analyse"] = _exporter_analyse(dossier, bilan_matrice)
    if "feuille" in formats:
        resultat["feuille"] = _exporter_feuille_de_route(dossier, sections)
    return resultat


# ---------------------------------------------------------------------------
# Texte des sections pour l'acheteur
# ---------------------------------------------------------------------------


def _texte_pour_acheteur(texte: str) -> str:
    """Retire les notes de travail : une citation `>` est une note de l'agent
    (critère servi, sources consultées), pas un contenu pour l'acheteur."""
    return "\n".join(l for l in texte.splitlines() if not l.lstrip().startswith(">"))


# ---------------------------------------------------------------------------
# HTML autonome — deux onglets + relevé des « [à compléter : …] »
# ---------------------------------------------------------------------------

_RE_A_COMPLETER = re.compile(r"\[\s*à compléter\s*:([^\]]*)\]")


def _exporter_html(dossier: Path, sections: list[dict]) -> Path:
    analyse = "\n\n".join(
        _lire_markdown(dossier / nom)
        for nom in ("02-analyse.md", "03-go-no-go.md")
    )
    onglets = [("Analyse", _vers_html(analyse))]
    parties_memoire: list[str] = []
    for section in sections:
        fichier = (dossier / "sections" / section["fichier"]) \
            if section["fichier"] else None
        if fichier is None or not fichier.exists():
            parties_memoire.append(
                f"<p class='manquante'>[section {section['numero']} manquante : "
                f"{html.escape(section['titre'])}]</p>"
            )
            continue
        parties_memoire.append(
            _vers_html(_texte_pour_acheteur(fichier.read_text(encoding="utf-8")))
        )
    onglets.append(("Mémoire", "\n".join(parties_memoire)))
    page = _page_html(onglets, _releve_a_completer(dossier, sections))
    chemin = dossier / "export" / "dossier.html"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(page, encoding="utf-8")
    return chemin


def _lire_markdown(chemin: Path) -> str:
    if not chemin.exists():
        return f"*(fichier {chemin.name} absent)*"
    return chemin.read_text(encoding="utf-8", errors="replace")


def _vers_html(texte: str) -> str:
    return charte.html_depuis_markdown(texte, ["tables"])


def _a_completer_par_fichier(
        dossier: Path, sections: list[dict]) -> list[tuple[str, list[str]]]:
    """Marqueurs `[à compléter : …]` relevés fichier par fichier, dans l'ordre."""
    fichiers = [dossier / "02-analyse.md", dossier / "03-go-no-go.md"]
    fichiers += [
        dossier / "sections" / s["fichier"] for s in sections if s["fichier"]
    ]
    groupes: list[tuple[str, list[str]]] = []
    for fichier in fichiers:
        if not fichier.exists():
            continue
        texte = fichier.read_text(encoding="utf-8", errors="replace")
        demandes = [
            t.group(1).strip() for t in _RE_A_COMPLETER.finditer(_sans_notes(texte))
            if t.group(1).strip() not in {"", "…", "..."}
        ]
        if demandes:
            groupes.append((fichier.name, demandes))
    return groupes


def _releve_a_completer(dossier: Path, sections: list[dict]) -> list[str]:
    groupes = _a_completer_par_fichier(dossier, sections)
    return [
        f"{nom} : [à compléter : {demande}]"
        for nom, demandes in groupes for demande in demandes
    ]


_SCRIPT_ONGLETS = """
function afficherOnglet(id) {
  document.querySelectorAll('.onglet').forEach(
    function (el) { el.hidden = el.id !== id; });
  document.querySelectorAll('nav button').forEach(function (b) {
    b.classList.toggle('actif', b.dataset.cible === id);
  });
}
document.querySelectorAll('nav button').forEach(function (b) {
  b.addEventListener('click', function () { afficherOnglet(b.dataset.cible); });
});
"""

_PAGE = """<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta http-equiv="Content-Security-Policy" content="{csp}">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Réponse — dossier de consultation</title>
<style>
  body {{ font-family: system-ui, -apple-system, sans-serif; margin: 0 auto;
         max-width: 60rem; padding: 0 1rem 3rem; color: #1c2430;
         background: #f7f8fa; line-height: 1.55; }}
  h1 {{ font-size: 1.4rem; margin: 1.5rem 0 .5rem; }}
  nav {{ display: flex; gap: .4rem; border-bottom: 2px solid #d7dce3;
        margin: 1rem 0 1.5rem; }}
  nav button {{ border: none; background: none; padding: .55rem .9rem;
        font-size: 1rem; cursor: pointer; color: #5a6472;
        border-bottom: 3px solid transparent; margin-bottom: -2px; }}
  nav button.actif {{ color: #103c8c; font-weight: 600;
        border-bottom-color: #103c8c; }}
  section.onglet {{ background: #fff; border: 1px solid #e3e7ec;
        border-radius: 8px; padding: 1.2rem 1.6rem; }}
  table {{ border-collapse: collapse; margin: .8rem 0; width: 100%; }}
  th, td {{ border: 1px solid #d7dce3; padding: .35rem .55rem;
        font-size: .95rem; text-align: left; vertical-align: top; }}
  th {{ background: #eef1f5; }}
  .manquante {{ color: #a33; }}
  aside {{ margin-top: 1.6rem; background: #fff7e8; border: 1px solid #ecd9ad;
        border-radius: 8px; padding: .8rem 1.2rem; }}
  aside ul {{ margin: .4rem 0 0; padding-left: 1.2rem; }}
  code {{ background: #eef1f5; border-radius: 4px; padding: .05rem .3rem; }}
</style>
</head>
<body>
<h1>Réponse à la consultation</h1>
<nav>
{boutons}
</nav>
{corps}
{releve}
<script>{script}</script>
</body>
</html>
"""


_CLASSE_ACTIF = ' class="actif"'


def _page_html(onglets: list[tuple[str, str]], releve: list[str]) -> str:
    boutons = "\n".join(
        f"<button type='button' data-cible='onglet-{index}'"
        f"{_CLASSE_ACTIF if index == 0 else ''}>"
        f"{html.escape(titre)}</button>"
        for index, (titre, _) in enumerate(onglets)
    )
    corps = "\n".join(
        f"<section class='onglet' id='onglet-{index}'"
        f"{' hidden' if index else ''}>{corps}</section>"
        for index, (_, corps) in enumerate(onglets)
    )
    if releve:
        bloc_releve = (
            "<aside><strong>Informations manquantes à fournir par "
            "l'entreprise</strong><ul>"
            + "".join(f"<li>{html.escape(ligne)}</li>" for ligne in releve)
            + "</ul></aside>"
        )
    else:
        bloc_releve = ""
    return _PAGE.format(boutons=boutons, corps=corps, releve=bloc_releve,
                        script=_SCRIPT_ONGLETS, csp=charte.csp(_SCRIPT_ONGLETS))


# ---------------------------------------------------------------------------
# Livrables de travail — analyse.html, feuille-de-route.html, matrice XLSX
# ---------------------------------------------------------------------------

_RE_TITRE_H2 = re.compile(r"^##\s+(.+?)\s*$", re.M)
_RE_MARQUEUR = re.compile(r"<!--\s*à remplir\s*-->")

LIBELLES_ETAPES = {
    "pieces": "Pièces", "analyse": "Analyse", "go-no-go": "Go / No-Go",
    "plan": "Plan", "redaction": "Rédaction", "relecture": "Relecture",
    "export": "Export",
}


def sections_h2(texte: str) -> list[tuple[str, str]]:
    """Découpe un markdown en (titre `##`, contenu sous le titre).

    Le texte avant le premier `##` est rendu sous titre vide ("").
    """
    titres = list(_RE_TITRE_H2.finditer(texte))
    if not titres:
        return [("", texte)]
    morceaux: list[tuple[str, str]] = []
    avant = texte[:titres[0].start()].strip()
    if avant:
        morceaux.append(("", avant))
    for rang, trouve in enumerate(titres):
        fin = titres[rang + 1].start() if rang + 1 < len(titres) else len(texte)
        morceaux.append((trouve.group(1).strip(), texte[trouve.end():fin].strip()))
    return morceaux


def _nettoyer_inline(texte: str) -> str:
    """Retire marqueur `<!-- à remplir -->` et emphase markdown d'une cellule."""
    texte = _RE_MARQUEUR.sub("", texte)
    for marqueur in ("**", "__", "`"):
        texte = texte.replace(marqueur, "")
    return texte.strip().strip("*").strip()


def _cellules(ligne: str) -> list[str]:
    """Cellules d'une ligne de tableau markdown (`| a | b |` → ["a", "b"])."""
    temporaire = ligne.strip().strip("|").replace("\\|", "\x00")
    return [c.strip().replace("\x00", "|") for c in temporaire.split("|")]


def _est_separateur(cellules: list[str]) -> bool:
    """Ligne `|---|:-:|` qui suit l'en-tête d'un tableau markdown."""
    return bool(cellules) and all(
        re.fullmatch(r":?-{2,}:?", cellule) for cellule in cellules
    )


def _lignes_tableau(texte: str) -> list[list[str]]:
    """Cellules de chaque ligne de tableau markdown d'un bloc, dans l'ordre."""
    return [_cellules(l) for l in texte.splitlines() if l.lstrip().startswith("|")]


def _contenu_section(texte: str, prefixe: str) -> str:
    """Contenu de la première section `##` dont le titre commence par `prefixe`."""
    for titre, contenu in sections_h2(texte):
        if titre.casefold().startswith(prefixe):
            return contenu
    return ""


def _cellule_a(cellules: list[str], indice: int | None) -> str:
    if indice is None or indice >= len(cellules):
        return ""
    return cellules[indice]


def _correspondance(dossier: Path) -> dict[str, tuple[str, str]]:
    """« Correspondance avec les critères » de 04-plan.md : numéro → (critère, poids)."""
    fichier_plan = dossier / "04-plan.md"
    if not fichier_plan.exists():
        return {}
    texte = fichier_plan.read_text(encoding="utf-8", errors="replace")
    for titre, contenu in sections_h2(texte):
        if "correspondance" not in titre.casefold():
            continue
        lignes = _lignes_tableau(contenu)
        if len(lignes) < 2 or not lignes[0][0].casefold().startswith("section"):
            return {}
        en_tetes = [cellule.casefold() for cellule in lignes[0]]
        col_critere = next(
            (i for i, e in enumerate(en_tetes) if e.startswith("critère")), None
        )
        col_poids = next((i for i, e in enumerate(en_tetes) if e == "poids"), None)
        correspondance: dict[str, tuple[str, str]] = {}
        for cellules in lignes[1:]:
            if _est_separateur(cellules):
                continue
            numero = re.match(r"\s*(\d+)", _nettoyer_inline(cellules[0]))
            if not numero:
                continue
            correspondance[numero.group(1).zfill(2)] = (
                _nettoyer_inline(_cellule_a(cellules, col_critere)),
                _nettoyer_inline(_cellule_a(cellules, col_poids)),
            )
        return correspondance
    return {}


# -- analyse.html et feuille-de-route.html ---------------------------------
# Les deux pages passent par la même fonction, charte.page_depuis_markdown :
# l'agent n'écrit que du markdown, la mise en page ne lui coûte rien.


def _sans_notes(texte: str) -> str:
    """Retire les notes de travail (citations `>`) et les marqueurs du gabarit."""
    lignes = [l for l in texte.splitlines() if not l.lstrip().startswith(">")]
    return _RE_MARQUEUR.sub("", "\n".join(lignes))


def _exporter_analyse(dossier: Path, bilan_matrice: str | None = None) -> Path:
    """analyse.html : l'analyse du DCE et le go/no-go, tels que rédigés, puis
    le renvoi vers la matrice de conformité quand elle existe."""
    morceaux = [
        _sans_notes(_lire_markdown(dossier / nom))
        for nom in ("02-analyse.md", "03-go-no-go.md")
    ]
    renvoi = _renvoi_matrice(dossier, bilan_matrice)
    if renvoi:
        morceaux.append(renvoi)
    return _ecrire_page(dossier, "analyse.html", "\n\n".join(morceaux), "Analyse du DCE")


def _renvoi_matrice(dossier: Path, bilan: str | None) -> str:
    """Paragraphe qui renvoie vers matrice-conformite.xlsx, rangée à côté de
    analyse.html. Sans bilan (matrice d'un export précédent), le lien seul."""
    if bilan is None and not (dossier / "export" / "matrice-conformite.xlsx").exists():
        return ""
    lien = "[matrice-conformite.xlsx](matrice-conformite.xlsx)"
    texte = f"Le détail est dans la matrice de conformité, {lien}"
    texte += f" : {bilan}." if bilan else "."
    return "## Matrice de conformité\n\n" + texte


def _date_limite(dossier: Path) -> str:
    """Valeur de la ligne « Date limite de remise » de 02-analyse.md, s'il y en a une."""
    texte = _lire_markdown(dossier / "02-analyse.md")
    for cellules in _lignes_tableau(texte):
        if cellules and cellules[0].casefold().startswith("date limite"):
            return _nettoyer_inline(_cellule_a(cellules, 1))
    return ""


def _feuille_de_route_markdown(dossier: Path, sections: list[dict]) -> str:
    """La feuille de route rédigée en markdown, à partir de l'état du dossier."""
    lignes = ["# Feuille de route", ""]
    date = _date_limite(dossier)
    if date:
        lignes += [f"**Date limite de remise :** {date}", ""]
    etat = espace.etat(dossier)
    libelles = {"faite": "faite", "en_cours": "en cours", "a_faire": "à faire"}
    lignes += ["## Avancement", "", "| Étape | État |", "|---|---|"]
    for etape, info in etat["etapes"].items():
        lignes.append(f"| {LIBELLES_ETAPES.get(etape, etape)} | {libelles[info['etat']]} |")
    correspondance = _correspondance(dossier)
    lignes += ["", "## Sections du mémoire", "",
               "| N° | Section | Rédigée | Critère noté | Poids |", "|---|---|---|---|---|"]
    for section in sections:
        critere, poids = correspondance.get(section["numero"], ("", ""))
        lignes.append(
            f"| {section['numero']} | {section['titre']} | "
            f"{'oui' if section['faite'] else 'non'} | {critere} | {poids} |"
        )
    groupes = _a_completer_par_fichier(dossier, sections)
    lignes += ["", "## À compléter", ""]
    if not groupes:
        lignes.append("Rien à compléter.")
    for nom, demandes in groupes:
        lignes += [f"**{nom}**", ""] + [f"- {demande}" for demande in demandes] + [""]
    pieces = _contenu_section(_lire_markdown(dossier / "02-analyse.md"), "pièces à remettre")
    if pieces:
        lignes += ["", "## Pièces à remettre", "", _sans_notes(pieces)]
    return "\n".join(lignes)


def _exporter_feuille_de_route(dossier: Path, sections: list[dict]) -> Path:
    texte = _feuille_de_route_markdown(dossier, sections)
    return _ecrire_page(dossier, "feuille-de-route.html", texte, "Feuille de route")


def _ecrire_page(dossier: Path, nom: str, texte: str, titre: str) -> Path:
    chemin = dossier / "export" / nom
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(charte.page_depuis_markdown(texte, titre), encoding="utf-8")
    return chemin


# -- matrice-conformite.xlsx ------------------------------------------------


def _cadre_du_dossier(dossier: Path) -> dict | None:
    """Résultat de cadre.extraire_cadre pour la pièce classée CRT du dossier.

    Rend None s'il n'y a pas d'inventaire ou pas exactement une pièce CRT.
    """
    fichier_etat = dossier / ".remporte" / "etat.json"
    if not fichier_etat.exists():
        return None
    donnees = json.loads(fichier_etat.read_text(encoding="utf-8"))
    candidats = [p for p in donnees.get("pieces", []) if p.get("type") == "CRT"]
    if len(candidats) != 1:
        return None
    return cadre.extraire_cadre(dossier / "dce" / candidats[0]["chemin"])


def _exporter_matrice(dossier: Path, sections: list[dict]) -> tuple[Path, str]:
    """matrice-conformite.xlsx : une ligne par exigence du CRT, sinon par section.

    Rend le chemin et le bilan en une phrase, repris par analyse.html.
    """
    from openpyxl import Workbook

    resultat = _cadre_du_dossier(dossier)
    entrees = cadre.couverture(dossier, resultat) if resultat else None
    if entrees is None and not sections:
        raise espace.Erreur(
            "Matrice vide : le CRT n'a pas d'exigences numérotées et le plan "
            "ne liste aucune section (remporte guide plan)."
        )
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Conformité"
    if entrees:
        _remplir_exigences(feuille, entrees, sections)
        restantes = sum(1 for entree in entrees if not entree["couverte"])
        bilan = (f"{_pluriel(len(entrees), 'exigence')} du cadre de réponse, "
                 + (f"dont {restantes} encore à traiter" if restantes else "toutes traitées"))
    else:
        _remplir_sections_plan(feuille, sections, _correspondance(dossier))
        restantes = sum(1 for section in sections if not section["faite"])
        bilan = (f"{_pluriel(len(sections), 'section')} du plan, "
                 + (f"dont {restantes} encore à rédiger" if restantes else "toutes rédigées"))
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions
    chemin = dossier / "export" / "matrice-conformite.xlsx"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(str(chemin))
    return chemin, bilan


def _pluriel(nombre: int, mot: str) -> str:
    return f"{nombre} {mot if nombre <= 1 else mot + 's'}"


def _remplir_exigences(
        feuille, entrees: list[dict], sections: list[dict]) -> None:
    """Une ligne par exigence numérotée du CRT, avec la section qui la cite."""
    labels = {
        s["fichier"]: f"{s['numero']} — {s['titre']}"
        for s in sections if s["fichier"]
    }
    feuille.append(
        ["Identifiant", "Chapitre", "Demande", "Section qui la cite", "Couverte"]
    )
    for entree in entrees:
        point = entree["point"]
        citante = " ; ".join(
            labels.get(nom, nom) for nom in entree["sections"]
        )
        feuille.append([
            entree["identifiant"] or "—",
            point.get("niveau") or "",
            point.get("demande") or point.get("titre") or "",
            citante or "—",
            "oui" if entree["couverte"] else "non",
        ])
    _mettre_en_forme(feuille, [13, 32, 72, 36, 11], (2, 3, 4))


def _remplir_sections_plan(
        feuille, sections: list[dict],
        correspondance: dict[str, tuple[str, str]]) -> None:
    """Sans exigences numérotées : une ligne par section du plan."""
    feuille.append(["Section", "Critère ou sous-critère noté", "Poids", "Rédigée"])
    for section in sections:
        critere, poids = correspondance.get(section["numero"], ("", ""))
        feuille.append([
            f"{section['numero']} — {section['titre']}", critere, poids,
            "oui" if section["faite"] else "non",
        ])
    _mettre_en_forme(feuille, [44, 44, 14, 11], (1, 2))


def _mettre_en_forme(
        feuille, largeurs: list[int], colonnes_wrappees: tuple[int, ...]) -> None:
    """En-tête en gras, largeurs lisibles, retour à la ligne sur le texte long.

    Aucune cellule n'est une formule : un texte du DCE qui commence par « = »
    (`=HYPERLINK(…)`, `=WEBSERVICE(…)`) reste du texte."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    for ligne in feuille.iter_rows():
        for cellule in ligne:
            if cellule.data_type == "f":
                cellule.data_type = "s"

    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for indice, largeur in enumerate(largeurs, 1):
        feuille.column_dimensions[get_column_letter(indice)].width = largeur
    for indice in colonnes_wrappees:
        for colonne in feuille.iter_cols(
                min_col=indice, max_col=indice, min_row=2):
            for cellule in colonne:
                cellule.alignment = Alignment(wrap_text=True, vertical="top")
