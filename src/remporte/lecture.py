"""Lecture des pièces du DCE : décompression d'archives et conversion en markdown.

Aucune fonction ne lève vers l'appelant : `convertir` rend toujours une
`Conversion`, même sur un fichier corrompu. Dépendances autorisées seulement :
pypdfium2 (PDF), python-docx (DOCX), openpyxl (XLSX), odfpy (ODS/ODT).
"""

from __future__ import annotations

import csv
import io
import re
import shutil
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

FORMATS_LUS = {
    ".pdf", ".docx", ".xlsx", ".xlsm", ".ods", ".odt", ".txt", ".md", ".csv",
    ".pptx",
}

# Seuil « PDF scanné » : moins de ce nombre de caractères extraits par page.
_SEUIL_SCAN_PAR_PAGE = 50

# Fichiers parasites des archives macOS, jamais des pièces du DCE.
_JUNK = {"__MACOSX", ".DS_Store"}

_ARCHIVES_NON_DECOMPRESSABLES = {".7z", ".rar"}


@dataclass
class Conversion:
    """Résultat de conversion d'une pièce en markdown.

    `texte` est vide dès que `statut` vaut autre chose que « ok ».
    """

    texte: str            # markdown ; "" si illisible
    statut: str           # "ok" | "vide" | "illisible" | "non_pris_en_charge"
    pages: int | None     # PDF uniquement
    motif: str | None     # explication si statut != "ok"


def _avertir(message: str) -> None:
    print(f"remporte : {message}", file=sys.stderr)


def decompresser(source: Path, cible: Path) -> list[Path]:
    """Copie un dossier ou décompresse un .zip (récursif) dans `cible`.

    Protège contre les chemins sortants (zip slip) et saute les fichiers
    parasites (__MACOSX, .DS_Store). Les .7z/.rar sont copiés tels quels et
    signalés sur stderr. Rend la liste des fichiers obtenus.
    """
    source = Path(source)
    cible = Path(cible)
    if source.is_dir():
        if cible.exists() and source.samefile(cible):
            return lister_fichiers(cible)  # copier sur soi-même : ne rien faire
        shutil.copytree(source, cible, dirs_exist_ok=True)
    elif zipfile.is_zipfile(source):
        cible.mkdir(parents=True, exist_ok=True)
        _extraire_zip(source, cible)
    else:
        # Fichier isolé (par exemple un .7z) : copié tel quel dans cible.
        cible.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, cible / source.name)
    _dezipper_recursif(cible)
    fichiers = lister_fichiers(cible)
    for fichier in fichiers:
        if fichier.suffix.lower() in _ARCHIVES_NON_DECOMPRESSABLES:
            _avertir(
                f"{fichier} : archive {fichier.suffix} copiée telle quelle "
                "(format non décompressable, utilisez 7-Zip puis relancez init)."
            )
    return fichiers


def lister_fichiers(base: Path) -> list[Path]:
    """Fichiers sous `base`, hors parasites, triés."""
    return [
        p for p in sorted(Path(base).rglob("*"), key=lambda p: p.as_posix())
        if p.is_file() and not _est_junk(p)
    ]


def _est_junk(chemin: Path) -> bool:
    return any(partie in _JUNK for partie in chemin.parts)


def _chemin_dangereux(nom: str) -> bool:
    """Vrai si une entrée d'archive tente de sortir de sa cible (zip slip)."""
    if not nom:
        return True
    pur = PurePosixPath(nom)
    if pur.is_absolute() or ".." in pur.parts:
        return True
    if re.match(r"^[A-Za-z]:", nom):  # lettre de lecteur Windows
        return True
    return False


def _extraire_zip(chemin_zip: Path, base: Path) -> None:
    """Extrait un .zip membre par membre, en sautant entrées dangereuses."""
    try:
        with zipfile.ZipFile(chemin_zip) as zf:
            base_resolue = base.resolve()
            for info in zf.infolist():
                nom = info.filename
                if info.is_dir() or _est_junk(Path(nom)):
                    continue
                if _chemin_dangereux(nom):
                    _avertir(f"entrée d'archive refusée (chemin sortant) : {nom}")
                    continue
                destination = base / nom
                if base_resolue not in destination.resolve().parents:
                    _avertir(f"entrée d'archive refusée (chemin sortant) : {nom}")
                    continue
                destination.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as source, open(destination, "wb") as sortie:
                    shutil.copyfileobj(source, sortie)
    except (zipfile.BadZipFile, OSError) as erreur:
        _avertir(f"{chemin_zip} : archive non extraite ({erreur})")


def _dezipper_recursif(base: Path) -> None:
    """Décompresse les .zip trouvés sous `base`, y compris zip dans zip."""
    traites: set[Path] = set()
    while True:
        zips = [
            p for p in sorted(base.rglob("*"), key=lambda p: p.as_posix())
            if p.is_file() and p.suffix.lower() == ".zip"
            and p.resolve() not in traites
        ]
        if not zips:
            return
        for zip_interne in zips:
            traites.add(zip_interne.resolve())
            destination = zip_interne.parent / (zip_interne.stem + "-dezippe")
            destination.mkdir(parents=True, exist_ok=True)
            _extraire_zip(zip_interne, destination)
            zip_interne.unlink(missing_ok=True)
            if not any(destination.rglob("*")):
                destination.rmdir()


def convertir(piece: Path) -> Conversion:
    """Convertit une pièce en markdown. Ne lève jamais."""
    piece = Path(piece)
    suffixe = piece.suffix.lower()
    try:
        if suffixe == ".pdf":
            return _convertir_pdf(piece)
        if suffixe == ".docx":
            return _convertir_docx(piece)
        if suffixe in {".xlsx", ".xlsm"}:
            return _convertir_xlsx(piece)
        if suffixe == ".ods":
            return _convertir_ods(piece)
        if suffixe == ".odt":
            return _convertir_odt(piece)
        if suffixe == ".pptx":
            return _convertir_pptx(piece)
        if suffixe == ".csv":
            return _convertir_csv(piece)
        if suffixe in {".txt", ".md"}:
            return _convertir_texte(piece)
        return Conversion(
            texte="",
            statut="non_pris_en_charge",
            pages=None,
            motif=f"format {suffixe or 'inconnu'} non pris en charge",
        )
    except Exception as erreur:  # noqa: BLE001 — la pièce est signalée, le traitement continue
        return Conversion(
            texte="",
            statut="illisible",
            pages=None,
            motif=f"{type(erreur).__name__}: {erreur}",
        )


# ---------------------------------------------------------------------------
# PDF (pypdfium2)
# ---------------------------------------------------------------------------

def _convertir_pdf(piece: Path) -> Conversion:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(piece))
    try:
        pages: list[str] = []
        for numero in range(len(pdf)):
            page = pdf[numero]
            texte = page.get_textpage().get_text_bounded()
            pages.append(texte.strip() if texte else "")
    finally:
        pdf.close()
    if not pages:
        return Conversion("", "vide", 0, "PDF sans page")
    caracteres = sum(len(texte) for texte in pages)
    if caracteres < _SEUIL_SCAN_PAR_PAGE * len(pages):
        return Conversion(
            "",
            "vide",
            len(pages),
            "PDF probablement scanné : texte non extractible",
        )
    parties = [pages[0]]
    for numero, texte in enumerate(pages[1:], start=2):
        parties.append(f"<!-- page {numero} -->")
        parties.append(texte)
    texte = "\n\n".join(parties).strip()
    texte = texte.replace("\r\n", "\n").replace("\r", "\n")
    return Conversion(texte, "ok", len(pages), None)


# ---------------------------------------------------------------------------
# DOCX (python-docx) — port de to_markdown.py, avec styles français « Titre N »
# ---------------------------------------------------------------------------

def _convertir_docx(piece: Path) -> Conversion:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(piece))
    paragraphes = {p._element: p for p in document.paragraphs}
    tableaux = {t._element: t for t in document.tables}
    lignes: list[str] = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraphe = paragraphes.get(element)
            if paragraphe is not None:
                md = _paragraphe_docx_vers_md(paragraphe)
                if md:
                    lignes.append(md)
        elif element.tag == qn("w:tbl"):
            table = tableaux.get(element)
            if table is not None:
                cellules = [_cellules_ligne_docx(ligne) for ligne in table.rows]
                md = _tableau_vers_md(cellules)
                if md:
                    lignes.append(md)
    texte = "\n\n".join(lignes)
    if not texte.strip():
        return Conversion("", "vide", None, "aucun texte dans le document")
    return Conversion(texte, "ok", None, None)


def _nom_style(paragraphe) -> str:
    try:
        return paragraphe.style.name if paragraphe.style else ""
    except Exception:  # noqa: BLE001 — style exotique : on retombe sur le texte brut
        return ""


def _niveau_titre_style(nom: str) -> int:
    nom = nom.strip().lower()
    trouve = re.match(r"^(?:heading|titre|title)\s*(\d+)$", nom)
    if trouve:
        return int(trouve.group(1))
    return 1 if nom in {"title", "titre"} else 0


def _est_puce(nom: str) -> bool:
    bas = nom.lower()
    return "bullet" in bas or "puce" in bas


def _est_numero(nom: str) -> bool:
    bas = nom.lower()
    return "number" in bas or "numérot" in bas or "numero" in bas


def _niveau_liste(nom: str) -> int:
    trouve = re.search(r"(\d+)\s*$", nom)
    return max(int(trouve.group(1)) - 1, 0) if trouve else 0


def _paragraphe_docx_vers_md(paragraphe) -> str:
    texte = paragraphe.text.strip()
    if not texte:
        return ""
    nom = _nom_style(paragraphe)
    niveau = _niveau_titre_style(nom)
    if niveau:
        return "#" * min(niveau, 6) + " " + texte
    if _est_puce(nom):
        return "  " * _niveau_liste(nom) + "- " + _runs_vers_md(paragraphe.runs)
    if _est_numero(nom):
        return "  " * _niveau_liste(nom) + "1. " + _runs_vers_md(paragraphe.runs)
    return _runs_vers_md(paragraphe.runs) or texte


def _runs_vers_md(runs) -> str:
    """Runs DOCX → markdown avec **gras** / *italique*, marqueurs fusionnés."""
    parties: list[str] = []
    for run in runs:
        if not run.text:
            continue
        gras = bool(run.font.bold)
        italique = bool(run.font.italic)
        if gras and italique:
            parties.append(f"***{run.text}***")
        elif gras:
            parties.append(f"**{run.text}**")
        elif italique:
            parties.append(f"*{run.text}*")
        else:
            parties.append(run.text)
    return _nettoyer_marqueurs("".join(parties))


def _nettoyer_marqueurs(texte: str) -> str:
    """Fusionne les marqueurs markdown de runs adjacents."""
    texte = re.sub(r"\*\*\*\*", " ", texte)
    texte = re.sub(r"(?<=\w)\*\*(?=\w)", " ", texte)
    return re.sub(r"  +", " ", texte)


def _cellules_ligne_docx(ligne) -> list[str]:
    return [" ".join(cellule.text.split()) for cellule in ligne.cells]


# ---------------------------------------------------------------------------
# Tableaux markdown partagés
# ---------------------------------------------------------------------------

def _cellule_md(texte: str) -> str:
    return texte.replace("|", "\\|").replace("\n", "<br>")


def _tableau_vers_md(lignes: list[list[str]]) -> str:
    """Lignes de cellules → tableau markdown (première ligne = en-tête)."""
    lignes = [ligne for ligne in lignes if ligne]
    if not lignes:
        return ""
    largeur = max(len(ligne) for ligne in lignes)
    lignes = [ligne + [""] * (largeur - len(ligne)) for ligne in lignes]
    sortie = ["| " + " | ".join(_cellule_md(c) for c in lignes[0]) + " |"]
    sortie.append("| " + " | ".join(["---"] * largeur) + " |")
    for ligne in lignes[1:]:
        sortie.append("| " + " | ".join(_cellule_md(c) for c in ligne) + " |")
    return "\n".join(sortie)


# ---------------------------------------------------------------------------
# XLSX / ODS — une section « ## Feuille <nom> » par feuille
# ---------------------------------------------------------------------------

def _convertir_xlsx(piece: Path) -> Conversion:
    from openpyxl import load_workbook

    # data_only=True : les formules non évaluées donnent leur valeur en cache.
    classeur = load_workbook(str(piece), read_only=True, data_only=True)
    try:
        sections: list[str] = []
        for feuille in classeur.worksheets:
            lignes: list[list[str]] = []
            for ligne in feuille.iter_rows(values_only=True):
                cellules = [
                    "" if valeur is None else str(valeur)
                    for valeur in ligne
                ]
                while cellules and not cellules[-1].strip():
                    cellules.pop()
                if any(cellule.strip() for cellule in cellules):
                    lignes.append([cellule.strip() for cellule in cellules])
            if lignes:
                sections.append(
                    f"## Feuille {feuille.title}\n\n" + _tableau_vers_md(lignes)
                )
    finally:
        classeur.close()
    texte = "\n\n".join(sections)
    if not texte.strip():
        return Conversion("", "vide", None, "classeur sans valeurs")
    return Conversion(texte, "ok", None, None)


# Garde-fou ODF : les vides de fin de feuille sont encodés par des
# répétitions énormes (≈ 1 048 576). On ne déplie jamais un vide ; seules
# les répétitions non vides atteignent le plafond.
_ODF_REPETITION_MAX = 1024


def _odf_repetition(element, attribut: str) -> int:
    brut = element.getAttribute(attribut)
    try:
        nombre = int(brut) if brut is not None else 1
    except (TypeError, ValueError):
        nombre = 1
    return max(1, min(nombre, _ODF_REPETITION_MAX))


def _lignes_feuille_ods(table) -> list[list[str]]:
    from odf import teletype
    from odf.table import TableCell, TableRow

    lignes: list[list[str]] = []
    for ligne in table.getElementsByType(TableRow):
        cellules: list[str] = []
        for cellule in ligne.getElementsByType(TableCell):
            texte = teletype.extractText(cellule).strip()
            repetition = _odf_repetition(cellule, "numbercolumnsrepeated")
            if not texte:
                repetition = 1
            cellules.extend([texte] * repetition)
        while cellules and not cellules[-1]:
            cellules.pop()
        repetition_ligne = _odf_repetition(ligne, "numberrowsrepeated")
        if not any(cellules):
            repetition_ligne = 1
        lignes.extend([list(cellules) for _ in range(repetition_ligne)])
    while lignes and not any(lignes[-1]):
        lignes.pop()
    return lignes


def _convertir_ods(piece: Path) -> Conversion:
    from odf.opendocument import load
    from odf.table import Table

    document = load(str(piece))
    sections: list[str] = []
    for table in document.getElementsByType(Table):
        lignes = [ligne for ligne in _lignes_feuille_ods(table) if any(ligne)]
        if lignes:
            nom = table.getAttribute("name") or "sans nom"
            sections.append(f"## Feuille {nom}\n\n" + _tableau_vers_md(lignes))
    texte = "\n\n".join(sections)
    if not texte.strip():
        return Conversion("", "vide", None, "classeur sans valeurs")
    return Conversion(texte, "ok", None, None)


def _convertir_odt(piece: Path) -> Conversion:
    from odf import teletype
    from odf.opendocument import load
    from odf.table import TableCell, TableRow

    document = load(str(piece))
    parties: list[str] = []
    for noeud in document.text.childNodes:
        local = getattr(noeud, "qname", (None, None))[1]
        if local == "table":
            lignes = []
            for ligne in noeud.getElementsByType(TableRow):
                cellules = [
                    teletype.extractText(cellule).strip()
                    for cellule in ligne.getElementsByType(TableCell)
                ]
                if any(cellules):
                    lignes.append(cellules)
            md = _tableau_vers_md(lignes)
            if md:
                parties.append(md)
        elif local == "h":
            texte = teletype.extractText(noeud).strip()
            if texte:
                try:
                    niveau = int(noeud.getAttribute("outlinelevel") or 1)
                except (TypeError, ValueError):
                    niveau = 1
                parties.append("#" * min(niveau, 6) + " " + texte)
        elif local == "p":
            texte = teletype.extractText(noeud).strip()
            if texte:
                parties.append(texte)
    texte = "\n\n".join(parties)
    if not texte.strip():
        return Conversion("", "vide", None, "aucun texte dans le document")
    return Conversion(texte, "ok", None, None)


# ---------------------------------------------------------------------------
# PPTX — lu en zip + XML (python-pptx hors dépendances autorisées)
# ---------------------------------------------------------------------------

_NS_A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"


def _numero_diapositive(nom: str) -> int:
    trouve = re.search(r"slide(\d+)\.xml$", nom)
    return int(trouve.group(1)) if trouve else 0


def _convertir_pptx(piece: Path) -> Conversion:
    sections: list[str] = []
    with zipfile.ZipFile(piece) as zf:
        diapositives = sorted(
            (nom for nom in zf.namelist()
             if re.fullmatch(r"ppt/slides/slide\d+\.xml", nom)),
            key=_numero_diapositive,
        )
        for nom in diapositives:
            racine = ElementTree.fromstring(zf.read(nom))
            textes = [
                element.text.strip()
                for element in racine.iter(_NS_A_T)
                if element.text and element.text.strip()
            ]
            if textes:
                sections.append(
                    f"## Diapositive {_numero_diapositive(nom)}\n\n"
                    + "\n\n".join(textes)
                )
    texte = "\n\n".join(sections)
    if not texte.strip():
        return Conversion("", "vide", None, "diaporama sans texte")
    return Conversion(texte, "ok", None, None)


# ---------------------------------------------------------------------------
# CSV / TXT / MD
# ---------------------------------------------------------------------------

def _lignes_csv(texte: str) -> list[list[str]]:
    """Parse un CSV (séparateur sniffé ; « ; » fréquent en France)."""
    echantillon = texte[:8192]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,\t|")
    except csv.Error:
        compteurs = {d: echantillon.count(d) for d in (";", ",", "\t", "|")}
        meilleur = max(compteurs, key=compteurs.get)
        delimiter = meilleur if compteurs[meilleur] > 0 else ","

        class DialecteRepli(csv.excel):
            pass

        DialecteRepli.delimiter = delimiter
        dialecte = DialecteRepli
    return [
        [" ".join((cellule or "").split()) for cellule in ligne]
        for ligne in csv.reader(io.StringIO(texte), dialecte)
    ]


def _convertir_csv(piece: Path) -> Conversion:
    brut = piece.read_bytes()
    try:
        texte = brut.decode("utf-8")
    except UnicodeDecodeError:
        texte = brut.decode("latin-1", errors="replace")
    lignes = [ligne for ligne in _lignes_csv(texte) if any(ligne)]
    if not lignes:
        return Conversion("", "vide", None, "CSV sans données")
    return Conversion(_tableau_vers_md(lignes), "ok", None, None)


def _convertir_texte(piece: Path) -> Conversion:
    texte = piece.read_text(encoding="utf-8", errors="replace").strip()
    if not texte:
        return Conversion("", "vide", None, "fichier vide")
    return Conversion(texte, "ok", None, None)
