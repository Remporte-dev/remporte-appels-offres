"""Tests des livrables de travail : analyse.html, feuille-de-route.html et
matrice-conformite.xlsx produits par `remporte exporter`."""

from __future__ import annotations

from pathlib import Path

import pytest
from conftest import fabriquer_docx

from remporte import espace, export

MOTS = " ".join(f"mot{index}" for index in range(60))

ANALYSE = """# Analyse du DCE

## Identité de la consultation

| | |
|---|---|
| Acheteur | Ville d'Orsay |
| Objet | Infogérance du SI municipal |
| Date limite de remise | 12 juin 2026 à 17 h 00 |

## Critères de jugement

| Critère | Pondération | Sous-critères et ce qui est noté | Source |
|---|---|---|---|
| Valeur technique | 50 % | Méthode et moyens | [RC art. 7] |
| Environnement | 10 % | Sobriété numérique | [RC art. 7] |
| Prix | 40 % | DPGF et DQE | [RC art. 7] |
| Cadre de réponse | [RC] | Trame imposée | [CRT] |

## Pièces à remettre

### Candidature
- DC1 et DC2.

### Offre
- Mémoire technique.

## Points de vigilance

- Pénalités de retard.
"""

GO_NO_GO = """# Go / No-Go

## Décision

**Go.** Le marché correspond au métier.

## Ce qu'il faut réunir pour répondre

- Deux références [à compléter : références comparables].
"""

PLAN = """# Plan du mémoire technique

## Sections

- [x] 01 — Moyens
- [ ] 02 — Prix

## Correspondance avec les critères

| Section | Critère ou sous-critère noté | Poids | Volume visé |
|---|---|---|---|
| 01 | Valeur technique | 50 % | 1 000 mots |
| 02 | Prix | 40 % | 300 mots |
"""


@pytest.fixture()
def dossier_pret(tmp_path: Path) -> Path:
    """Dossier de réponse rempli : analyse, go/no-go, plan, une section rédigée."""
    source = tmp_path / "DCE"
    source.mkdir()
    fabriquer_docx(source / "RC.docx")
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    (dossier / "02-analyse.md").write_text(ANALYSE, encoding="utf-8")
    (dossier / "03-go-no-go.md").write_text(GO_NO_GO, encoding="utf-8")
    (dossier / "04-plan.md").write_text(PLAN, encoding="utf-8")
    (dossier / "sections" / "01-moyens.md").write_text(
        f"# Moyens\n\n{MOTS}. La conformité PP 001 est traitée ici.\n",
        encoding="utf-8",
    )
    return dossier


def test_analyse_html(dossier_pret: Path):
    resultat = export.exporter(dossier_pret, {"analyse"})
    assert resultat["analyse"] == dossier_pret / "export" / "analyse.html"
    page = resultat["analyse"].read_text(encoding="utf-8")
    assert "Infogérance du SI municipal" in page  # objet en en-tête
    assert "Ville d'Orsay" in page  # acheteur
    assert "12 juin 2026 à 17 h 00" in page  # date limite de remise
    assert "50 %" in page and "40 %" in page  # critères pondérés
    assert "<strong>Go.</strong>" in page  # verdict
    assert "Méthode :" not in page  # notes de travail du gabarit retirées
    assert "<style>" in page and "</html>" in page  # page autonome complète
    assert _aucune_ressource_externe(page)
    assert 'class="brand"' in page  # logo de la charte
    assert "remporte.fr" not in page and "utm_" not in page  # aucun contenu promotionnel


def test_feuille_de_route_html(dossier_pret: Path):
    resultat = export.exporter(dossier_pret, {"feuille"})
    assert resultat["feuille"] == dossier_pret / "export" / "feuille-de-route.html"
    page = resultat["feuille"].read_text(encoding="utf-8")
    assert "Feuille de route" in page
    assert "12 juin 2026 à 17 h 00" in page  # date limite rappelée
    assert "Rédaction" in page and "en cours" in page  # étapes
    assert "<td>01</td>" in page and "Moyens" in page
    assert "Valeur technique" in page and "50 %" in page  # correspondance
    assert "<td>02</td>" in page and "non" in page
    assert "<li>…</li>" not in page  # exemple du gabarit non relevé
    assert "références comparables" in page  # à compléter
    assert "03-go-no-go.md" in page  # regroupement par fichier
    assert "DC1 et DC2" in page  # pièces à remettre
    assert _aucune_ressource_externe(page)


def test_matrice_exigences_crt(dossier_pret: Path, monkeypatch):
    """Avec exigences numérotées : identifiant, chapitre, demande, citation."""
    points = [
        {"numero": 1, "titre": "PP 001 — Le candidat décrit ses moyens.",
         "demande": "PP 001 — Le candidat décrit ses moyens.",
         "niveau": "2.1 : Moyens"},
        {"numero": 2, "titre": "PP 002 — Le candidat justifie ses prix.",
         "demande": "PP 002 — Le candidat justifie ses prix.",
         "niveau": "4 : Prix"},
    ]
    monkeypatch.setattr(export, "_cadre_du_dossier",
                        lambda dossier: {"origine": "liste", "points": points})
    resultat = export.exporter(dossier_pret, {"matrice"})
    assert resultat["matrice"] == \
        dossier_pret / "export" / "matrice-conformite.xlsx"

    from openpyxl import load_workbook

    feuille = load_workbook(str(resultat["matrice"])).active
    assert [cellule.value for cellule in feuille[1]] == [
        "Identifiant", "Chapitre", "Demande", "Section qui la cite", "Couverte"]
    assert feuille.max_row == 3
    assert [cellule.value for cellule in feuille[2]] == [
        "PP 001", "2.1 : Moyens", "PP 001 — Le candidat décrit ses moyens.",
        "01 — Moyens", "oui"]
    assert [cellule.value for cellule in feuille[3]][4] == "non"
    assert feuille.freeze_panes == "A2"
    assert feuille.auto_filter.ref
    assert feuille.column_dimensions["C"].width > 0  # largeurs lisibles


def test_matrice_sections_sans_crt(dossier_pret: Path):
    """Sans exigences numérotées : une ligne par section du plan."""
    resultat = export.exporter(dossier_pret, {"matrice"})
    from openpyxl import load_workbook

    feuille = load_workbook(str(resultat["matrice"])).active
    assert [cellule.value for cellule in feuille[1]] == [
        "Section", "Critère ou sous-critère noté", "Poids", "Rédigée"]
    assert [cellule.value for cellule in feuille[2]] == [
        "01 — Moyens", "Valeur technique", "50 %", "oui"]
    assert [cellule.value for cellule in feuille[3]] == [
        "02 — Prix", "Prix", "40 %", "non"]
    assert feuille.freeze_panes == "A2"
    assert feuille.auto_filter.ref


def test_matrice_vide_refusee(tmp_path: Path):
    """Ni exigences numérotées, ni section au plan : erreur explicite."""
    source = tmp_path / "DCE"
    source.mkdir()
    fabriquer_docx(source / "RC.docx")
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    (dossier / "04-plan.md").write_text(
        "# Plan\n\n## Sections\n\nRien encore.\n", encoding="utf-8")
    with pytest.raises(espace.Erreur):
        export.exporter(dossier, {"matrice"})


def test_cli_exporter_produit_tout(dossier_pret: Path, capsys):
    """`remporte exporter` sans --formats produit les quatre fichiers."""
    from remporte import cli

    assert cli.main(["exporter", "--dossier", str(dossier_pret)]) == 0
    sortie = capsys.readouterr().out
    for cle in ("html", "analyse", "feuille", "matrice"):
        assert f"{cle} : " in sortie
    for nom in ("dossier.html", "analyse.html",
                "feuille-de-route.html", "matrice-conformite.xlsx"):
        assert (dossier_pret / "export" / nom).exists()


def test_couverture_partagee_cli_et_export(dossier_pret: Path, capsys):
    """La couverture de `remporte cadre` et la matrice partagent le calcul."""
    from remporte import cadre, cli

    points = [
        {"numero": 1, "titre": "PP 001 — Le candidat décrit ses moyens.",
         "demande": "PP 001 — Le candidat décrit ses moyens.", "niveau": ""},
    ]
    resultat = {"origine": "liste", "points": points}
    entrees = cadre.couverture(dossier_pret, resultat)
    assert len(entrees) == 1
    assert entrees[0]["couverte"] and entrees[0]["sections"] == ["01-moyens.md"]
    assert cadre.couverture(dossier_pret, {"origine": "aucune", "points": []}) \
        is None
    code = cli.main([
        "cadre", "--couverture", "--dossier", str(dossier_pret),
    ])
    sortie = capsys.readouterr()
    assert code == 1  # sans pièce CRT : message d'usage, pas de crash
    assert "Aucun CRT identifié" in sortie.out


def test_commande_html(tmp_path):
    """`remporte html` : une page à la charte, sans ressource réseau, depuis du markdown."""
    from remporte import cli

    source = tmp_path / "note.md"
    source.write_text("# Note de synthèse\n\n| Critère | Poids |\n|---|---|\n| Prix | 40 % |\n",
                      encoding="utf-8")
    assert cli.main(["html", str(source)]) == 0
    page = (tmp_path / "note.html").read_text(encoding="utf-8")
    assert "<title>Note de synthèse</title>" in page and "<td>40 %</td>" in page
    assert "#1d2b50" in page and _aucune_ressource_externe(page)
    assert "remporte.fr" not in page and "utm_" not in page  # aucun contenu promotionnel


def _aucune_ressource_externe(page: str) -> bool:
    """Page autonome : aucune image, feuille de style ou script chargé du réseau."""
    return not any(motif in page for motif in ('src="http', "<link", "url(http", "@import"))


def test_analyse_renvoie_vers_la_matrice(dossier_pret: Path):
    """Exportées ensemble, l'analyse renvoie vers la matrice avec son bilan."""
    resultat = export.exporter(dossier_pret, {"analyse", "matrice"})
    page = resultat["analyse"].read_text(encoding="utf-8")
    assert "Matrice de conformité" in page
    assert 'href="matrice-conformite.xlsx"' in page
    assert "2 sections du plan, dont 1 encore à rédiger" in page


def test_analyse_sans_matrice_ni_renvoi(dossier_pret: Path):
    """Sans matrice exportée, l'analyse ne renvoie vers rien."""
    page = export.exporter(dossier_pret, {"analyse"})["analyse"].read_text(encoding="utf-8")
    assert "matrice-conformite.xlsx" not in page


def test_analyse_seule_garde_le_lien_vers_une_matrice_existante(dossier_pret: Path):
    """Une matrice d'un export précédent reste citée, sans bilan recalculé."""
    export.exporter(dossier_pret, {"matrice"})
    page = export.exporter(dossier_pret, {"analyse"})["analyse"].read_text(encoding="utf-8")
    assert 'href="matrice-conformite.xlsx"' in page
    assert "encore à rédiger" not in page
