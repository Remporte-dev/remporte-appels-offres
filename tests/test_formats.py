"""Tests de remporte.formats : détection des formats attendus par l'acheteur.

Cas positifs (chaque motif) et négatifs (bordereaux de prix exclus, DCE sans
contrainte), section « Formats attendus » de 01-pieces.md et commande
`remporte formats`.
"""

from __future__ import annotations

import json
from pathlib import Path

from conftest import ecrire_pdf, fabriquer_docx, fabriquer_xlsx

from remporte import espace, formats
from remporte.cli import main


def docx_avec_texte(chemin: Path, paragraphes: list[str]) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Règlement de consultation", 1)
    for paragraphe in paragraphes:
        document.add_paragraph(paragraphe)
    document.save(str(chemin))
    return chemin


def xlsx_avec_texte(chemin: Path, lignes: list[list[str]]) -> Path:
    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    for ligne in lignes:
        feuille.append(ligne)
    classeur.save(str(chemin))
    return chemin


def initialiser(tmp_path: Path, pieces: dict[str, Path]) -> Path:
    source = tmp_path / "DCE"
    source.mkdir()
    for nom, piece in pieces.items():
        (source / nom).write_bytes(piece.read_bytes())
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    return dossier


def test_soutenance_et_limite_pages(tmp_path: Path):
    rc = docx_avec_texte(tmp_path / "1-RC.docx", [
        "Les candidats short-listés seront auditionnés par la commission.",
        "Le mémoire technique est limité à 30 pages maximum.",
    ])
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    constats = formats.detecter(dossier)
    types = [(c["type"], c.get("pages")) for c in constats]
    assert ("soutenance", None) in types
    assert ("limite_pages", 30) in types
    soutenance = next(c for c in constats if c["type"] == "soutenance")
    assert soutenance["piece"] == "1-RC.docx"
    assert "audition" in soutenance["extrait"]
    assert len(soutenance["extrait"]) <= 200
    pages = next(c for c in constats if c["type"] == "limite_pages")
    assert "30 pages maximum" in pages["extrait"]


def test_variantes_de_limite_de_pages(tmp_path: Path):
    rc = docx_avec_texte(tmp_path / "1-RC.docx", [
        "La réponse ne doit pas dépasser 40 pages.",
        "Chaque section est limitée à 5 pages.",
        "Un maximum de 60 pages est admis pour l'ensemble.",
    ])
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    pages = sorted(
        c["pages"] for c in formats.detecter(dossier)
        if c["type"] == "limite_pages"
    )
    assert pages == [5, 40, 60]


def test_pas_de_faux_positif_sur_une_page_40(tmp_path: Path):
    rc = docx_avec_texte(tmp_path / "1-RC.docx", [
        "Voir page 40 du règlement pour le plan de masse.",
    ])
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    assert formats.detecter(dossier) == []


def test_excel_impose_et_bordereaux_exclus(tmp_path: Path):
    crt = xlsx_avec_texte(tmp_path / "2-CRT.xlsx", [
        ["Exigence", "Réponse du candidat"], ["PP 001", ""],
    ])
    bpu = fabriquer_xlsx(tmp_path / "3-BPU.xlsx")
    annexe = xlsx_avec_texte(tmp_path / "4-Annexe-technique.xlsx", [
        ["Cadre à compléter par le candidat"], [""],
    ])
    annexe_muette = xlsx_avec_texte(tmp_path / "5-Annexe-plans.xlsx", [
        ["Plan", "Révision"], ["A", "1"],
    ])
    dossier = initialiser(tmp_path, {
        "2-CRT.xlsx": crt,
        "3-BPU.xlsx": bpu,
        "4-Annexe-technique.xlsx": annexe,
        "5-Annexe-plans.xlsx": annexe_muette,
    })
    excel = [c for c in formats.detecter(dossier) if c["type"] == "excel_impose"]
    assert [c["piece"] for c in excel] == ["2-CRT.xlsx", "4-Annexe-technique.xlsx"]


def test_cadre_word(tmp_path: Path):
    crt = docx_avec_texte(tmp_path / "2-Cadre-de-reponse.docx", [
        "Le candidat répond point par point dans le présent cahier.",
    ])
    dossier = initialiser(tmp_path, {"2-Cadre-de-reponse.docx": crt})
    cadre = [c for c in formats.detecter(dossier) if c["type"] == "cadre_word"]
    assert [c["piece"] for c in cadre] == ["2-Cadre-de-reponse.docx"]
    assert "cahier" in cadre[0]["extrait"].lower()


def test_dce_sans_contrainte(tmp_path: Path):
    rc = fabriquer_docx(tmp_path / "1-RC.docx")
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    assert formats.detecter(dossier) == []


def test_section_formats_dans_01_pieces(tmp_path: Path):
    rc = docx_avec_texte(tmp_path / "1-RC.docx", [
        "Une audition est prévue ; le mémoire est limité à 30 pages maximum.",
    ])
    crt = fabriquer_xlsx(tmp_path / "2-CRT.xlsx")
    dossier = initialiser(tmp_path, {"1-RC.docx": rc, "2-CRT.xlsx": crt})
    section = (dossier / "01-pieces.md").read_text(encoding="utf-8")
    assert "## Formats attendus" in section
    assert "Soutenance ou présentation orale" in section
    assert "Limite de pages (30 pages)" in section
    assert "`dce/1-RC.docx`" in section
    assert "Classeur Excel à compléter" in section
    assert "audition" in section


def test_section_formats_vide(tmp_path: Path):
    rc = fabriquer_docx(tmp_path / "1-RC.docx")
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    section = (dossier / "01-pieces.md").read_text(encoding="utf-8")
    assert "Aucune contrainte de forme détectée." in section


def test_commande_formats(tmp_path: Path, capsys):
    rc = docx_avec_texte(tmp_path / "1-RC.docx", [
        "Support de présentation demandé pour la présentation orale.",
    ])
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    assert main(["formats", "--dossier", str(dossier)]) == 0
    sortie = capsys.readouterr().out
    assert "Soutenance ou présentation orale" in sortie
    assert "dce/1-RC.docx" in sortie

    assert main(["formats", "--dossier", str(dossier), "--json"]) == 0
    constats = json.loads(capsys.readouterr().out)["constats"]
    assert {c["type"] for c in constats} == {"soutenance"}


def test_commande_formats_sans_constat(tmp_path: Path, capsys):
    rc = fabriquer_docx(tmp_path / "1-RC.docx")
    dossier = initialiser(tmp_path, {"1-RC.docx": rc})
    assert main(["formats", "--dossier", str(dossier)]) == 0
    assert "Aucune contrainte de forme détectée" in capsys.readouterr().out


def test_soutenance_ignoree_dans_le_cctp(tmp_path):
    """Une présentation orale prévue pendant l'exécution (CCTP) n'est pas une audition."""
    from docx import Document

    from remporte import espace, formats

    source = tmp_path / "DCE"
    source.mkdir()
    document = Document()
    document.add_paragraph("Cahier des clauses techniques particulières")
    document.add_paragraph("Le titulaire fait une présentation orale au comité de pilotage.")
    document.save(source / "CCTP.docx")
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    assert not [c for c in formats.detecter(dossier) if c["type"] == "soutenance"]
