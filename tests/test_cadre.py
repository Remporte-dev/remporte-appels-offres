"""Tests de remporte.cadre : styles dédiés puis motifs, seuil des 3 points."""

from __future__ import annotations

from pathlib import Path

from conftest import fabriquer_docx_crt, ecrire_pdf

from remporte.cadre import extraire_cadre


def test_cadre_par_styles(tmp_path: Path):
    piece = fabriquer_docx_crt(tmp_path / "CRT.docx")
    resultat = extraire_cadre(piece)
    assert resultat["origine"] == "styles"
    assert len(resultat["points"]) == 3
    premier = resultat["points"][0]
    assert set(premier) == {"numero", "titre", "demande", "niveau"}
    assert premier["numero"] == 1
    assert premier["demande"] == "Le candidat décrit son organisation."
    assert premier["niveau"] == "Cahier des réponses"


def test_cadre_par_motifs_sur_texte(tmp_path: Path):
    piece = tmp_path / "CRT.txt"
    piece.write_text(
        "# Présentation\n"
        "Le candidat décrit son organisation.\n"
        "Texte neutre sans demande.\n"
        "Le candidat fournit trois références.\n"
        "Le candidat justifie ses moyens.\n",
        encoding="utf-8",
    )
    resultat = extraire_cadre(piece)
    assert resultat["origine"] == "motifs"
    assert len(resultat["points"]) == 3
    assert resultat["points"][0]["niveau"] == "Présentation"


def test_cadre_motifs_sur_pdf(tmp_path: Path):
    piece = ecrire_pdf(tmp_path / "crt.pdf", [
        "Presentation du cadre de reponse"
        "\nLe candidat decrit ses references."
        "\nLe candidat fournit un planning."
        "\nLe candidat indique ses moyens.",
    ])
    resultat = extraire_cadre(piece)
    assert resultat["origine"] == "motifs"
    assert len(resultat["points"]) == 3


def test_cadre_insuffisant_est_aucune(tmp_path: Path):
    piece = tmp_path / "faible.txt"
    piece.write_text(
        "Le candidat décrit son organisation.\nRien d'autre ici.\n",
        encoding="utf-8",
    )
    resultat = extraire_cadre(piece)
    assert resultat == {"origine": "aucune", "points": []}


def test_cadre_piece_scannee_est_aucune(tmp_path: Path):
    piece = ecrire_pdf(tmp_path / "scan.pdf", ["", ""])
    assert extraire_cadre(piece) == {"origine": "aucune", "points": []}


def test_liste_exigences_numerotees(tmp_path: Path):
    """Une liste « {PP 001.} chapitre : demande  page » fait foi, toutes tournures."""
    from docx import Document

    document = Document()
    document.add_paragraph("Liste des points à préciser")
    demandes = [
        "Le candidat précise l'organisation.",
        "Le candidat propose des indicateurs.",
        "Le candiat détaille la réversibilité.",
        "Afin de tenir les délais, le candidat propose une mobilisation.",
        "Le candidat décrit la méthode.",
    ]
    for rang, demande in enumerate(demandes, 1):
        document.add_paragraph(f"{{PP {rang:03d}.}}\t2.{rang} : Chapitre {rang} : {demande}\t{rang + 7}")
    chemin = tmp_path / "CRT.docx"
    document.save(chemin)
    cadre = extraire_cadre(chemin)
    assert cadre["origine"] == "liste"
    assert len(cadre["points"]) == 5
    assert cadre["points"][2]["titre"].startswith("PP 003 — Le candiat détaille")
    assert cadre["points"][0]["niveau"] == "2.1 : Chapitre 1"
