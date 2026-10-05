"""Tests de remporte.export : DOCX assemblé, HTML autonome à deux onglets."""

from __future__ import annotations

from pathlib import Path

import pytest
from conftest import fabriquer_docx

from remporte import espace, export


@pytest.fixture()
def dossier_pret(tmp_path: Path) -> Path:
    """Dossier de réponse avec plan complet et deux sections rédigées."""
    source = tmp_path / "DCE"
    source.mkdir()
    fabriquer_docx(source / "RC.docx")
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    (dossier / "04-plan.md").write_text(
        "# Plan\n\n"
        "- [x] 01 — Moyens humains\n"
        "- [x] 02 — Références\n",
        encoding="utf-8",
    )
    (dossier / "sections/01-moyens-humains.md").write_text(
        "# Moyens humains\n\n"
        "L'équipe compte **trois** ingénieurs et un chef de projet.\n\n"
        "| Profil | Effectif |\n|---|---|\n| Ingénieur | 3 |\n\n"
        "- Certification ISO 27001\n- Habilitation secret\n",
        encoding="utf-8",
    )
    (dossier / "sections/02-references.md").write_text(
        "# Références\n\n"
        "Trois références comparables, dont une [à compléter : effectif du marché].\n",
        encoding="utf-8",
    )
    return dossier


def test_exporter_refuse_plan_vide(tmp_path: Path):
    source = tmp_path / "DCE"
    source.mkdir()
    fabriquer_docx(source / "RC.docx")
    dossier = tmp_path / "reponse"
    espace.initialiser(source, dossier)
    with pytest.raises(espace.Erreur):
        export.exporter(dossier, {"docx"})


def test_exporter_docx(tmp_path: Path, dossier_pret: Path):
    resultat = export.exporter(dossier_pret, {"docx"})
    chemin = resultat["docx"]
    assert chemin == dossier_pret / "export" / "memoire.docx"
    assert chemin.exists()

    from docx import Document

    document = Document(str(chemin))
    textes = [p.text for p in document.paragraphs]
    assert "Mémoire technique" in textes
    assert "01 — Moyens humains" in textes
    assert any("trois" in t for t in textes)
    assert any(t.style.name.startswith("List Bullet") for t in document.paragraphs)
    tableau = document.tables[0]
    assert tableau.cell(0, 0).text == "Profil"
    assert tableau.cell(1, 1).text == "3"


def test_exporter_html(tmp_path: Path, dossier_pret: Path):
    resultat = export.exporter(dossier_pret, {"html"})
    chemin = resultat["html"]
    assert chemin == dossier_pret / "export" / "dossier.html"
    page = chemin.read_text(encoding="utf-8")
    assert "<style>" in page and "<script>" in page  # autonome
    assert "Analyse" in page and "Mémoire" in page
    assert "afficherOnglet" in page
    assert "<table>" in page  # tableau markdown rendu
    assert "à compléter : effectif du marché" in page  # relevé
    assert "http://" not in page and "https://" not in page  # zéro ressource


def test_exporter_les_deux_formats(tmp_path: Path, dossier_pret: Path):
    resultat = export.exporter(dossier_pret, {"docx", "html"})
    assert resultat["docx"].exists() and resultat["html"].exists()


def test_exporter_section_manquante(tmp_path: Path, dossier_pret: Path):
    (dossier_pret / "sections/02-references.md").unlink()
    resultat = export.exporter(dossier_pret, {"docx"})
    from docx import Document

    document = Document(str(resultat["docx"]))
    assert any("section 02 manquante" in p.text for p in document.paragraphs)


def test_export_sans_doublon_ni_notes(tmp_path: Path, dossier_pret: Path):
    """Titre propre à la section, notes `>` retirées, lignes coupées recollées."""
    from docx import Document

    plan = dossier_pret / "04-plan.md"
    plan.write_text("- [x] 01 — Moyens\n", encoding="utf-8")
    for ancien in (dossier_pret / "sections").glob("*.md"):
        ancien.unlink()
    (dossier_pret / "sections" / "01-moyens.md").write_text(
        "# 01 — Moyens\n\n> Sert le critère 1.2 [RC art. 7].\n\n"
        "Une équipe de trois\ningénieurs dédiés.\n",
        encoding="utf-8",
    )
    export.exporter(dossier_pret, {"docx", "html"})
    textes = [p.text for p in Document(dossier_pret / "export" / "memoire.docx").paragraphs]
    assert textes.count("01 — Moyens") == 1
    assert "Une équipe de trois ingénieurs dédiés." in textes
    assert not any("Sert le critère" in t for t in textes)
    page = (dossier_pret / "export" / "dossier.html").read_text(encoding="utf-8")
    assert "Sert le critère" not in page
