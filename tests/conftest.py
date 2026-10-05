"""Fabriques de pièces de test : PDF minimal en bytes, DOCX, XLSX, ODS, PPTX, ZIP."""

from __future__ import annotations

import zipfile
from pathlib import Path


# ---------------------------------------------------------------------------
# PDF minimal — construit à la main (objets + xref exacts), lu par pypdfium2
# ---------------------------------------------------------------------------

def ecrire_pdf(chemin: Path, pages: list[str]) -> Path:
    """Petit PDF valide : pages de lignes Helvetica (texte ASCII ; \n = ligne).

    Numérotation : 1 = catalogue, 2 = arbre des pages, 3 = fonte,
    4..3+n = pages, 4+n..3+2n = contenus.
    """
    nb = len(pages)
    objets: dict[int, bytes] = {}
    objets[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{3 + i} 0 R" for i in range(1, nb + 1))
    objets[2] = f"<< /Type /Pages /Kids [{kids}] /Count {nb} >>".encode()
    objets[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i in range(1, nb + 1):
        objets[3 + i] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {3 + nb + i} 0 R >>"
        ).encode()
    for i, page in enumerate(pages, start=1):
        commandes = []
        for k, ligne in enumerate(page.split("\n")):
            echappe = (ligne.replace("\\", r"\\")
                       .replace("(", r"\(").replace(")", r"\)"))
            y = 720 - 20 * k
            commandes.append(f"1 0 0 1 72 {y} Tm ({echappe}) Tj")
        flux = (f"BT /F1 14 Tf " + " ".join(commandes) + " ET").encode("latin-1")
        objets[3 + nb + i] = (
            f"<< /Length {len(flux)} >>\nstream\n".encode() + flux + b"\nendstream"
        )

    sortie = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for numero in sorted(objets):
        offsets[numero] = len(sortie)
        sortie += f"{numero} 0 obj\n".encode() + objets[numero] + b"\nendobj\n"
    depart_xref = len(sortie)
    maxi = max(objets) + 1
    sortie += f"xref\n0 {maxi}\n".encode()
    sortie += b"0000000000 65535 f \n"
    for numero in range(1, maxi):
        sortie += f"{offsets[numero]:010d} 00000 n \n".encode()
    sortie += (
        f"trailer\n<< /Size {maxi} /Root 1 0 R >>\n"
        f"startxref\n{depart_xref}\n%%EOF\n"
    ).encode()
    chemin.write_bytes(bytes(sortie))
    return chemin


# ---------------------------------------------------------------------------
# DOCX / XLSX / ODS / PPTX / ZIP
# ---------------------------------------------------------------------------

def fabriquer_docx(chemin: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Règlement de consultation", 1)
    document.add_paragraph("Objet : AMO pour le musée.")
    document.add_paragraph("Critères de jugement :", style="Normal")
    document.add_paragraph(
        "Le critère valeur technique est noté sur 60 %.", style="List Bullet"
    )
    document.add_paragraph("Le candidat **doit** justifier ses références.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Critère"
    table.cell(0, 1).text = "Pondération"
    table.cell(1, 0).text = "Valeur technique"
    table.cell(1, 1).text = "60 %"
    document.save(str(chemin))
    return chemin


def fabriquer_docx_crt(chemin: Path) -> Path:
    """CRT avec un style dédié « DemandeCandidat » pour la voie par styles."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    document.add_heading("Cahier des réponses", 1)
    style = document.styles.add_style("DemandeCandidat", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph(
        "Le candidat décrit son organisation.", style="DemandeCandidat"
    )
    document.add_paragraph(
        "Le candidat fournit trois références comparables.",
        style="DemandeCandidat",
    )
    document.add_paragraph(
        "Le candidat justifie ses moyens humains.", style="DemandeCandidat"
    )
    document.add_paragraph("Présentation générale de l'offre.")
    document.save(str(chemin))
    return chemin


def fabriquer_xlsx(chemin: Path) -> Path:
    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "BPU"
    feuille.append(["Désignation", "Quantité", "Prix"])
    feuille.append(["Prestation A", 3, None])
    feuille2 = classeur.create_sheet("DQE")
    feuille2.append(["Poste", "Montant"])
    feuille2.append(["Lot 1", 1200])
    classeur.save(str(chemin))
    return chemin


def fabriquer_ods(chemin: Path) -> Path:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    table = Table(name="Feuille1")
    for ligne in (["Désignation", "Prix"], ["Prestation X", ""]):
        tr = TableRow()
        for valeur in ligne:
            tc = TableCell(valuetype="string")
            tc.addElement(P(text=valeur))
            tr.addElement(tc)
        table.addElement(tr)
    document.spreadsheet.addElement(table)
    document.save(str(chemin))
    return chemin


def fabriquer_pptx(chemin: Path) -> Path:
    ns = (
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    )
    diapo = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f"<p:sld {ns}><p:cSld><p:spTree>"
        "<a:p><a:r><a:t>Présentation de l'entreprise</a:t></a:r></a:p>"
        "<a:p><a:r><a:t>Créée en 2010</a:t></a:r></a:p>"
        "</p:spTree></p:cSld></p:sld>"
    )
    with zipfile.ZipFile(chemin, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("ppt/slides/slide1.xml", diapo)
    return chemin


def fabriquer_zip(chemin: Path, contenus: dict[str, bytes | str],
                  zips_internes: dict[str, dict[str, str]] | None = None) -> Path:
    """ZIP avec fichiers plats et, si demandé, des ZIP dans le ZIP."""
    with zipfile.ZipFile(chemin, "w") as zf:
        for nom, contenu in contenus.items():
            if isinstance(contenu, str):
                contenu = contenu.encode("utf-8")
            zf.writestr(nom, contenu)
        for nom_zip, fichiers in (zips_internes or {}).items():
            tampon = Path(str(chemin) + ".interne.zip")
            with zipfile.ZipFile(tampon, "w") as zi:
                for nom, contenu in fichiers.items():
                    zi.writestr(nom, contenu.encode("utf-8"))
            zf.writestr(nom_zip, tampon.read_bytes())
            tampon.unlink()
    return chemin
